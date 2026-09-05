"""
Build viva presentation (.pptx) following the template style.
Compact 14 slides with no redundant content.
"""
import os
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

TEMPLATE = "/Users/jasmi/Downloads/Presentation Template May Viva.pptx"
OUTPUT   = "/Users/jasmi/Desktop/AV-Deepfake1M/Try/presentation.pptx"
PROJECT  = "/Users/jasmi/Desktop/AV-Deepfake1M/Try"

BLUE  = RGBColor(0x00, 0x70, 0xC0)
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK  = RGBColor(0x1A, 0x1A, 0x2E)
FONT  = "Helvetica Neue"

prs = Presentation(TEMPLATE)

# Remove template slides 2-6
for idx in range(len(prs.slides) - 1, 0, -1):
    rId = prs.slides._sldIdLst[idx].get(qn('r:id'))
    prs.part.drop_rel(rId)
    prs.slides._sldIdLst.remove(prs.slides._sldIdLst[idx])

LAYOUT_TITLE   = 0   # Title & Subtitle
LAYOUT_CONTENT = 13  # Title and Content

def find(slide, name):
    for s in slide.shapes:
        if s.name == name: return s
    return None

def title(slide, text):
    t = find(slide, "Title 1")
    if t:
        t.text_frame.clear()
        r = t.text_frame.paragraphs[0].add_run()
        r.text = text

def content(slide, lines):
    """lines = list of (text, bold, level). level 0 = normal, 1 = indented."""
    c = find(slide, "Content Placeholder 2")
    if not c: return
    tf = c.text_frame; tf.clear()
    for i, (text, bold, level) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = level
        r = p.add_run(); r.text = text; r.font.name = FONT
        if bold: r.font.bold = True

def make(title_text, body_lines):
    s = prs.slides.add_slide(prs.slide_layouts[LAYOUT_CONTENT])
    title(s, title_text)
    content(s, body_lines)
    return s

def table_slide(slide_title, headers, rows, col_widths, body_lines=()):
    s = prs.slides.add_slide(prs.slide_layouts[LAYOUT_CONTENT])
    title(s, slide_title)
    n_rows = len(rows) + 1; n_cols = len(headers)
    ts = s.shapes.add_table(n_rows, n_cols, Inches(0.5), Inches(1.4), Inches(12.3), Inches(0.8 + 0.5 * len(rows)))
    t = ts.table
    if col_widths:
        for ci, w in enumerate(col_widths):
            t.columns[ci].width = w
    for ci, h in enumerate(headers):
        cell = t.cell(0, ci); cell.text = h
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs: r.font.bold = True; r.font.size = Pt(12); r.font.color.rgb = WHITE; r.font.name = FONT
        tcPr = cell._tc.get_or_add_tcPr()
        sf = cell._tc.makeelement(qn('a:solidFill'), {})
        sc = sf.makeelement(qn('a:srgbClr'), {'val': '0070C0'})
        sf.append(sc); tcPr.append(sf)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.cell(ri + 1, ci); cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for r in p.runs: r.font.size = Pt(11); r.font.name = FONT; r.font.color.rgb = BLACK
    if body_lines:
        y = Inches(1.4) + Inches(0.5 * (len(rows) + 1))
        tb = s.shapes.add_textbox(Inches(0.5), y, Inches(12.3), Inches(6.5) - y)
        tf = tb.text_frame; tf.clear()
        for i, (text, bold, level) in enumerate(body_lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.level = level
            r = p.add_run(); r.text = text; r.font.name = FONT; r.font.size = Pt(14)
            if bold: r.font.bold = True
    return s

def image_slide(slide_title, img_path, extra_lines=()):
    s = prs.slides.add_slide(prs.slide_layouts[LAYOUT_CONTENT])
    title(s, slide_title)
    if os.path.exists(img_path):
        s.shapes.add_picture(img_path, Inches(0.5), Inches(1.5), Inches(7.0))
    if extra_lines:
        tb = s.shapes.add_textbox(Inches(8.0), Inches(1.5), Inches(4.8), Inches(5.5))
        tf = tb.text_frame; tf.clear()
        for i, (text, bold, level) in enumerate(extra_lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.level = level
            r = p.add_run(); r.text = text; r.font.name = FONT; r.font.size = Pt(14)
            if bold: r.font.bold = True
    return s


# === SLIDE 1: Title (edit existing) ===
s1 = prs.slides[0]
t = find(s1, "Title 1")
t.text_frame.clear()
p = t.text_frame.paragraphs[0]
r1 = p.add_run(); r1.text = "Viva: "; r1.font.size = Pt(66); r1.font.bold = True; r1.font.color.rgb = BLUE; r1.font.name = FONT
r2 = p.add_run(); r2.text = "Deepfake Detection Using\nCross-Model Transformer Fusion"; r2.font.size = Pt(48); r2.font.bold = True; r2.font.color.rgb = BLACK; r2.font.name = FONT
sub = find(s1, "Subtitle 2")
sub.text_frame.clear()
for i, line in enumerate(["Jasmi Preethi Alasapuri", "2571395", "", "Supervisor: Lucian Duta", "", "BSc (Hons) Data Science and Artificial Intelligence | CN6000 Dissertation 2026"]):
    p = sub.text_frame.paragraphs[0] if i == 0 else sub.text_frame.add_paragraph()
    r = p.add_run(); r.text = line; r.font.name = FONT; r.font.size = Pt(22); r.font.color.rgb = BLACK


# === SLIDE 2: The Problem ===
make("The Problem — Why Deepfake Detection Matters", [
    ("AI-generated synthetic media is eroding trust in digital evidence across law, politics, finance, and journalism.", False, 0),
    ("", False, 0),
    ("Real-world incidents:", True, 0),
    ("• £20M corporate fraud via deepfake video call (Arup, 2024)", False, 0),
    ("• Political deepfakes for disinformation (Zelenskyy, 2022)", False, 0),
    ("• AI kidnapping and impersonation scams (FBI, 2025)", False, 0),
    ("", False, 0),
    ("Multimodal deepfakes — manipulating both audio and video — are the hardest to detect.", False, 0),
    ("Current systems are vision-centric and fail to capture speech-lip temporal relationships.", False, 0),
    ("", False, 0),
    ("Computing challenge: Process 1.4 TB video data · Train 50M-parameter network · Real-time inference", True, 0),
])

# === SLIDE 3: Aims & Objectives ===
make("Project Aims & Objectives", [
    ("Aim: A multimodal deepfake detection system to distinguish real from manipulated audio-visual media using deep learning.", True, 0),
    ("", False, 0),
    ("Six Objectives:", True, 0),
    ("", False, 0),
    ("1. Literature review on deepfake generation and detection techniques", False, 0),
    ("2. Analyse real-world impacts; identify gaps in current solutions", False, 0),
    ("3. Quantitative secondary analysis of AV-Deepfake1M++ (speaker-disjoint split)", False, 0),
    ("4. Design and implement Cross-Modal Transformer Fusion architecture", False, 0),
    ("5. Resumable training pipeline; evaluation on held-out test set", False, 0),
    ("6. Standalone inference system and web interface for non-technical use", False, 0),
])

# === SLIDE 4: Literature Review ===
table_slide("Literature Review — Key Findings",
    ["Phase", "Era", "Characteristics"],
    [["Visual Fidelity", "2017–2019", "Face-swaps via Autoencoders / GANs"],
     ["In-the-Wild", "2019–2021", "Multi-subject, uncontrolled environments"],
     ["Multimodal", "2023–2025", "Neural generation: audio + video"]],
    [Inches(2.5), Inches(2.2), Inches(7.6)],
    body_lines=[
        ("", False, 0),
        ("Three Gaps Identified:", True, 0),
        ("1. Identity Leakage — Random splits allow face recognition, not manipulation detection", False, 0),
        ("2. Vision-Centric Bias — Audio is handled via simple concatenation or ignored entirely", False, 0),
        ("3. Limited Generalisation — Detectors overfit to training generators and datasets", False, 0),
    ])

# === SLIDE 5: Architecture ===
s5 = prs.slides.add_slide(prs.slide_layouts[LAYOUT_CONTENT])
title(s5, "Cross-Modal Transformer Fusion — Architecture")
if os.path.exists(os.path.join(PROJECT, "figures/architecture.png")):
    s5.shapes.add_picture(os.path.join(PROJECT, "figures/architecture.png"),
                          Inches(0.3), Inches(1.3), width=Inches(12.7))

# === SLIDE 6: Design Decisions ===
table_slide("Key Design Decisions",
    ["Component", "Initial Plan", "Final Choice", "Reason"],
    [["Audio Encoder", "Wav2Vec 2.0", "ResNet18 + Mel-spec", "torchaudio/FFmpeg handles corrupted MP4s"],
     ["Video Encoder", "MobileNetV3", "ResNet3D-18", "3D convolutions for temporal lip-motion artefacts"],
     ["Fusion Module", "DiMoDif", "Transformer + [CLS]", "Cross-modal self-attention captures inconsistencies"],
     ["Loss Function", "BCE", "Focal Loss (γ=2.0)", "Down-weights easy examples; focuses on hard cases"]],
    [Inches(2.0), Inches(2.0), Inches(3.0), Inches(5.3)])

# === SLIDE 7: Implementation ===
s7 = prs.slides.add_slide(prs.slide_layouts[LAYOUT_CONTENT])
title(s7, "Implementation — Training Pipeline")
tp = os.path.join(PROJECT, "figures/two_phase_training.png")
if os.path.exists(tp):
    s7.shapes.add_picture(tp, Inches(0.3), Inches(1.4), width=Inches(6.8))
# Add table on this slide:
ts = s7.shapes.add_table(9, 2, Inches(7.5), Inches(1.4), Inches(5.3), Inches(5.5))
t = ts.table
t.columns[0].width = Inches(2.0); t.columns[1].width = Inches(3.3)
data = [["Setting", "Value"],
        ["Dataset", "68,851 clips (val split); 4 manipulation types"],
        ["Phase 1", "Frozen encoders (~2 epochs)"],
        ["Phase 2", "Full fine-tune at 10× lower LR"],
        ["Optimiser", "AdamW (1e-4 fusion, 1e-5 encoders)"],
        ["Batch Size", "8 per GPU (×2 DataParallel)"],
        ["Loss", "Focal Loss (γ=2.0, α=0.25)"],
        ["Infrastructure", "Vast.ai (2× RTX 3080); W&B tracking"],
        ["Workers", "28 CPU parallel extraction, crash-resumable"]]
for ri, row in enumerate(data):
    for ci, val in enumerate(row):
        cell = t.cell(ri, ci); cell.text = val
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size = Pt(11); r.font.name = FONT
                if ri == 0: r.font.bold = True; r.font.color.rgb = WHITE
        if ri == 0:
            tcPr = cell._tc.get_or_add_tcPr()
            sf = cell._tc.makeelement(qn('a:solidFill'), {})
            sc = sf.makeelement(qn('a:srgbClr'), {'val': '0070C0'})
            sf.append(sc); tcPr.append(sf)

# === SLIDE 8: Results ===
s8 = prs.slides.add_slide(prs.slide_layouts[LAYOUT_CONTENT])
title(s8, "Results — Training History & Per-Type Accuracy")
th = os.path.join(PROJECT, "comparison_results/training_history.png")
pt = os.path.join(PROJECT, "figures/per_type_accuracy_bar_chart.png")
if os.path.exists(th): s8.shapes.add_picture(th, Inches(0.1), Inches(1.3), height=Inches(2.8))
if os.path.exists(pt): s8.shapes.add_picture(pt, Inches(6.3), Inches(1.3), height=Inches(2.8))

ts = s8.shapes.add_table(5, 6, Inches(0.5), Inches(4.4), Inches(12.3), Inches(2.2))
t = ts.table
for ci, w in enumerate([Inches(1.2), Inches(0.8), Inches(1.8), Inches(1.8), Inches(1.8), Inches(4.9)]):
    t.columns[ci].width = w
hdrs = ["Model", "Ep", "Val AUC", "Test AUC", "Test Acc", "Notes"]
rows = [["M1", "1", "0.663", "—", "—", "Underfit; corrupted download"],
        ["M2", "5", "0.994", "0.919", "66%", "Highest val AUC"],
        ["M3 ★", "3", "0.985", "0.937", "93%", "Best: Precision 1.000, 0 false positives"],
        ["M4", "5", "0.993", "0.915", "71%", "Same session as M3, epoch 5"]]
for ri, row in enumerate([hdrs] + rows):
    for ci, val in enumerate(row):
        cell = t.cell(ri, ci); cell.text = val
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size = Pt(10); r.font.name = FONT
                if ri == 0: r.font.bold = True; r.font.color.rgb = WHITE
        if ri == 0:
            tcPr = cell._tc.get_or_add_tcPr(); sf = cell._tc.makeelement(qn('a:solidFill'), {})
            sc = sf.makeelement(qn('a:srgbClr'), {'val': '0070C0'})
            sf.append(sc); tcPr.append(sf)

# === SLIDE 9: Model Comparison & Calibration ===
s9 = prs.slides.add_slide(prs.slide_layouts[LAYOUT_CONTENT])
title(s9, "Model Comparison & Score Calibration")
mc = os.path.join(PROJECT, "comparison_results/model_comparison.png")
cc = os.path.join(PROJECT, "figures/calibration_curves.png")
if os.path.exists(mc): s9.shapes.add_picture(mc, Inches(0.1), Inches(1.3), height=Inches(2.8))
if os.path.exists(cc): s9.shapes.add_picture(cc, Inches(6.3), Inches(1.3), height=Inches(2.8))
tb = s9.shapes.add_textbox(Inches(0.5), Inches(4.3), Inches(12.3), Inches(3.0))
tf = tb.text_frame
for text, bold, _ in [("Key Finding:", True, 0),
    ("Higher validation AUC does NOT equal better deployment performance.", True, 0),
    ("", False, 0),
    ("• Model 2 (AUC 0.994) → 66% accuracy at τ=0.5 (poorly calibrated near boundary)", False, 0),
    ("• Model 3 (AUC 0.985) → 93% accuracy at τ=0.5 (well-separated, zero false positives)", False, 0),
    ("", False, 0),
    ("Calibration curves confirm Model 3 has near-perfect score-probability alignment. Audio-modified hardest type (88%) — genuine video signal counteracts audio head.", False, 0)]:
    p = tf.paragraphs[0] if tf.paragraphs[0].text == '' else tf.add_paragraph()
    r = p.add_run(); r.text = text; r.font.name = FONT; r.font.size = Pt(14)
    if bold: r.font.bold = True

# === SLIDE 10: Key Findings ===
make("Key Findings", [
    ("1. Phase 2 fine-tuning is essential", True, 0),
    ("   Phase 1 only (Model 1): AUC 0.663. Phase-2 models: ≥ 0.915 test AUC.", False, 0),
    ("", False, 0),
    ("2. Validation AUC ≠ deployment performance", True, 0),
    ("   Model 2 (AUC 0.994) → 66%. Model 3 (AUC 0.985) → 93%. Score calibration matters.", False, 0),
    ("", False, 0),
    ("3. Three-head architecture provides genuine modality specialisation", True, 0),
    ("   Audio/video scores dissociate by manipulation type as predicted. Per-modality interpretability.", False, 0),
    ("", False, 0),
    ("4. Speaker-disjoint evaluation is critical", True, 0),
    ("   GroupShuffleSplit ensures zero speaker overlap. Results reflect manipulation detection, not identity.", False, 0),
])

# === SLIDE 11: Web Interface ===
s11 = prs.slides.add_slide(prs.slide_layouts[LAYOUT_CONTENT])
title(s11, "Web Interface — Three Tabs")
imgs = [
    ("figures/web_analyze_fake.png", "Analyze — Upload + Verdict", Inches(0.5), Inches(1.5)),
    ("figures/web_compare.png", "Compare — Side-by-Side Models", Inches(4.7), Inches(1.5)),
    ("figures/web_history.png", "History — SQLite-Backed Log", Inches(8.9), Inches(1.5)),
]
for path, label, left, top in imgs:
    full = os.path.join(PROJECT, path)
    if os.path.exists(full):
        s11.shapes.add_picture(full, left, top, width=Inches(3.8))
    lbl = s11.shapes.add_textbox(left, Inches(5.0), Inches(3.8), Inches(0.4))
    p = lbl.text_frame.paragraphs[0]; r = p.add_run(); r.text = label
    r.font.name = FONT; r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = BLUE
    p.alignment = PP_ALIGN.CENTER
tb = s11.shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(12.3), Inches(0.8))
p = tb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Flask backend + inference.py CLI. Drag-and-drop, model selector, per-head scores (audio/video/joint), batch processing, PDF reports."
r.font.name = FONT; r.font.size = Pt(13); r.font.color.rgb = BLACK

# === SLIDE 12: Limitations ===
table_slide("Limitations",
    ["Limitation", "Impact"],
    [["100-video test set", "1 misclassification = 1% accuracy shift; wide confidence intervals"],
     ["Validation split only", "68,851 clips vs 1M+ in full training set; limited diversity"],
     ["No ablation study", "Cannot attribute performance to Focal Loss vs BCE, Transformer vs MLP"],
     ["Fixed 2-second window", "May miss manipulation in short or late segments"],
     ["Single dataset", "No cross-dataset evaluation (FakeAVCeleb, DFDC, FaceForensics++)"],
     ["5-epoch training cap", "Cloud GPU budget prevented full convergence and hyperparameter sweeps"]],
    [Inches(3.5), Inches(8.8)])

# === SLIDE 13: Future Work ===
make("Future Work", [
    ("1. Ablation studies — Isolate Focal Loss vs BCE, Transformer vs MLP, full-frame vs lip-region", False, 0),
    ("2. Cross-dataset evaluation — Test on FakeAVCeleb, DFDC, FaceForensics++ for generalisation", False, 0),
    ("3. Full dataset training — Complete AV-Deepfake1M++ (1M+ clips) instead of validation split", False, 0),
    ("4. Temporal localisation — Frame-level predictions via fake_segments annotations", False, 0),
    ("5. Threshold optimisation — Calibrate on held-out set to balance FP/FN for deployment", False, 0),
    ("6. Longer training runs — 20+ epochs with hyperparameter sweeps (institutional GPU)", False, 0),
])

# === SLIDE 14: Conclusion ===
make("Conclusion", [
    ("What Was Achieved:", True, 0),
    ("• Multimodal deepfake detector: 93% accuracy, precision 1.000, zero false positives", False, 0),
    ("• Six objectives met — scope exceeds initial CN6000 proposal", False, 0),
    ("• Speaker-disjoint evaluation for honest generalisation estimates", False, 0),
    ("• Interpretable three-head output revealing per-modality vulnerability patterns", False, 0),
    ("• Production-ready: Resumable pipeline, W&B audit trail, web UI, standalone CLI", False, 0),
    ("", False, 0),
    ("Key Takeaway:", True, 0),
    ("Score calibration, not AUC, determines deployment readiness. Model 3 (epoch 3, lower AUC) outperformed Model 2 (epoch 5) by 27pp — with zero false positives.", False, 0),
    ("", False, 0),
    ("Hardest challenges were engineering: corrupted MP4s, crash-resumable extraction, cloud GPU management.", False, 0),
])

prs.save(OUTPUT)
print(f"Saved {OUTPUT} — {len(prs.slides)} slides")
